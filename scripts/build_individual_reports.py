#!/usr/bin/env python3
"""Build five polished Finfolio personal reports from Individual Footprint files."""

from __future__ import annotations

import re
import shutil
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REFERENCE = Path(
    "/Users/admin/Documents/TCH442_Group12/fomc-price-discovery/output/pdf/"
    "Group12_FOMC_Report_Reference_Style.docx"
)
SOURCE_DIR = ROOT / "docs" / "individual" / "en"
OUTPUT_DIR = ROOT / "formatted_reports"

COURSE = "Information Technology Applications in Finance and Banking (NHA408E)"
INSTRUCTOR = "Phan Trần Trung Dũng"
REPORT_DATE = "Hanoi, October 2026"
PROJECT = "FINFOLIO 2.0 — PORTFOLIO ATTRIBUTION & RISK ANALYTICS"

MEMBERS = {
    "linh": {
        "name": "Hoàng Khánh Linh",
        "student_id": "2412380024",
        "role": "Product Owner / Financial Logic Lead",
        "toc_pages": [3, 3, 3, 4, 5, 5, 6],
        "abstract": (
            "This report documents my responsibility for product scope and the financial methodology of Finfolio 2.0. "
            "It focuses on return attribution, Euler risk attribution, FX attribution, optimization assumptions and "
            "analytical claim boundaries, with traceable evidence in the team repository."
        ),
    },
    "quynh-anh": {
        "name": "Nguyễn Quỳnh Anh",
        "student_id": "2413380010",
        "role": "Product Manager / Flow & Testing Lead",
        "toc_pages": [3, 3, 3, 4, 5, 5, 6],
        "abstract": (
            "This report documents my responsibility for converting product requirements into user flows, acceptance "
            "criteria and verifiable tests. It covers happy, alternative and error paths, regression evidence, bug "
            "governance and production acceptance."
        ),
    },
    "bao-an": {
        "name": "Lê Bảo An",
        "student_id": "2412380002",
        "role": "Team Representative / Lead Developer",
        "toc_pages": [3, 3, 4, 5, 5, 6, 7],
        "abstract": (
            "This report documents my responsibility for architecture, implementation and integration in Finfolio 2.0. "
            "It covers the financial engine, constrained optimization, complete Monte Carlo workflow, Web Worker, "
            "automated verification, CI/CD and production deployment."
        ),
    },
    "minh-ngoc": {
        "name": "Trần Minh Ngọc",
        "student_id": "2412380036",
        "role": "Business Analyst / Data Lead",
        "toc_pages": [3, 3, 3, 4, 5, 5, 6],
        "abstract": (
            "This report documents my responsibility for data readiness and provenance in Finfolio 2.0. It covers the "
            "input dictionary, source register, benchmark and fallback policy, currency and common-date alignment, "
            "instrument classification, validation and freshness monitoring."
        ),
    },
    "ngoc-anh": {
        "name": "Nguyễn Ngọc Anh",
        "student_id": "2413380008",
        "role": "UI/UX Designer / Visual Design Lead",
        "toc_pages": [3, 3, 3, 4, 5, 5, 6],
        "abstract": (
            "This report documents my responsibility for information architecture and user experience in Finfolio 2.0. "
            "It covers the Forest/Paper design system, Vietnamese-compatible typography, multi-page navigation, "
            "native-SVG visualization, responsive layout and progressive disclosure."
        ),
    },
}

BLUE = "1F4E79"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "EEF5FA"
MID_GREY = "666666"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="B8C7D1", size="4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def add_page_number(paragraph) -> None:
    for child in list(paragraph._p):
        paragraph._p.remove(child)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, sep, end])


def add_toc(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Mở file trong Microsoft Word và chọn Update Field để cập nhật mục lục."
    separate.append(placeholder)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


def add_static_toc(doc: Document, pages: list[int]) -> None:
    titles = [
        "I. Role and accountability",
        "II. Distinctive contribution to Finfolio 2.0",
        "III. Work completed across the project",
        "IV. Evidence and reproducibility",
        "V. Weekly evidence trail",
        "VI. Limitations and lessons learned",
        "Individual declaration",
    ]
    entries = list(zip(titles, map(str, pages)))
    for title, page in entries:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.right_indent = Cm(0.5)
        p.paragraph_format.space_after = Pt(5)
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "right")
        tab.set(qn("w:leader"), "dot")
        tab.set(qn("w:pos"), "8700")
        tabs.append(tab)
        p._p.get_or_add_pPr().append(tabs)
        run = p.add_run(f"{title}\t{page}")
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)


def add_hyperlink(paragraph, text: str, url: str):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    run_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_pr.extend([color, underline])
    run.append(run_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE = re.compile(
    r"(\[([^\]]+)\]\(([^)]+)\)|\*\*([^*]+)\*\*|`([^`]+)`|\$([^$]+)\$)"
)


def add_inline(paragraph, text: str) -> None:
    pos = 0
    for match in INLINE.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos : match.start()])
        token = match.group(0)
        if token.startswith("["):
            add_hyperlink(paragraph, match.group(2), match.group(3))
        elif token.startswith("**"):
            paragraph.add_run(match.group(4)).bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(match.group(5))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(70, 70, 70)
        else:
            math_text = match.group(6)
            for old, new in (
                (r"\mathbf{w}", "w"),
                (r"\sigma", "σ"),
                (r"\Sigma", "Σ"),
                (r"\sum", "∑"),
                (r"\sqrt", "√"),
                (r"\cdot", "·"),
                (r"\to", "→"),
                (r"\%", "%"),
                ("^T", "ᵀ"),
            ):
                math_text = math_text.replace(old, new)
            math_text = math_text.replace("{", "(").replace("}", ")")
            math_text = math_text.replace("\\", "")
            run = paragraph.add_run(math_text)
            run.font.name = "Cambria Math"
            run.italic = True
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def clear_document(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for style_name, size, color, before, after in (
        ("Title", 25, BLUE, 0, 10),
        ("Subtitle", 13, MID_GREY, 0, 8),
        ("Heading 1", 16, BLUE, 16, 6),
        ("Heading 2", 13, "2F75B5", 10, 4),
        ("Heading 3", 11, "2F75B5", 8, 3),
    ):
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        if style_name.startswith("Heading"):
            style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Report Abstract" not in styles:
        abstract = styles.add_style("Report Abstract", WD_STYLE_TYPE.PARAGRAPH)
    else:
        abstract = styles["Report Abstract"]
    abstract.font.name = "Times New Roman"
    abstract.font.size = Pt(11)
    abstract.font.italic = False
    abstract.paragraph_format.line_spacing = 1.35
    abstract.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    abstract.paragraph_format.left_indent = Cm(0.6)
    abstract.paragraph_format.right_indent = Cm(0.6)
    abstract.paragraph_format.space_after = Pt(6)

    for style_name in ("List Bullet", "List Number"):
        if style_name not in styles:
            style = styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
            style.base_style = styles["Normal"]
        else:
            style = styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(11)
        style.paragraph_format.line_spacing = 1.25
        style.paragraph_format.space_after = Pt(3)


def add_cover_border(section) -> None:
    sect_pr = section._sectPr
    borders = sect_pr.find(qn("w:pgBorders"))
    if borders is None:
        borders = OxmlElement("w:pgBorders")
        sect_pr.append(borders)
    borders.set(qn("w:display"), "firstPage")
    borders.set(qn("w:offsetFrom"), "page")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "double")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "24")
        node.set(qn("w:color"), "000000")
        borders.append(node)


def configure_cover_section(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7875)
    section.bottom_margin = Inches(0.7875)
    section.left_margin = Inches(1.18125)
    section.right_margin = Inches(0.7875)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)
    section.different_first_page_header_footer = True
    add_cover_border(section)


def configure_body_section(doc: Document):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.5)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)
    section.different_first_page_header_footer = False
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    section.header.paragraphs[0].text = ""
    footer_p = section.footer.paragraphs[0]
    add_page_number(footer_p)
    for run in footer_p.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(10)

    sect_pr = section._sectPr
    inherited_borders = sect_pr.find(qn("w:pgBorders"))
    if inherited_borders is not None:
        sect_pr.remove(inherited_borders)
    page_num = sect_pr.find(qn("w:pgNumType"))
    if page_num is None:
        page_num = OxmlElement("w:pgNumType")
        sect_pr.append(page_num)
    page_num.set(qn("w:start"), "2")
    return section


def add_cover(doc: Document, meta: dict, logo_path: Path) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("FOREIGN TRADE UNIVERSITY")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor.from_string(BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("FACULTY OF BANKING AND FINANCE")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    p.add_run().add_picture(str(logo_path), width=Cm(2.5))

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(10)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("PERSONAL PROJECT REPORT")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(24)
    r.font.color.rgb = RGBColor.from_string(BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run(PROJECT)
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(17)

    table = doc.add_table(rows=6, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(4.2)
    table.columns[1].width = Cm(10.4)
    fields = [
        ("Student", meta["name"]),
        ("Student ID", meta["student_id"]),
        ("Role", meta["role"]),
        ("Group", "G09 — K63"),
        ("Course", COURSE),
        ("Instructor", INSTRUCTOR),
    ]
    for row, (label, value) in zip(table.rows, fields):
        row.cells[0].text = label
        row.cells[1].text = value
        set_cell_shading(row.cells[0], LIGHT_BLUE)
        for idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, 110, 120, 110, 120)
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(10.5)
                    run.bold = idx == 0

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(32)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(REPORT_DATE)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)


def add_front_matter(doc: Document, meta: dict) -> None:
    heading = doc.add_paragraph("ABSTRACT", style="Reference Front")
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(meta["abstract"], style="Report Abstract")

    p = doc.add_paragraph(style="Report Abstract")
    r = p.add_run("Keywords: ")
    r.bold = True
    p.add_run("Finfolio 2.0; portfolio attribution; risk analytics; individual contribution; reproducibility.")

    doc.add_paragraph("Verification principle", style="Heading 2")
    p = doc.add_paragraph(style="Report Abstract")
    add_inline(
        p,
        "This report interprets the member's responsibilities from the G09 Individual Footprint. "
        "Each contribution is linked to an output, repository file or reproducible check. Git commit "
        "authorship is not treated as sole proof because one member may integrate files for the team.",
    )

    doc.add_paragraph("CONTENTS", style="Reference Front").alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_static_toc(doc, meta["toc_pages"])
    doc.add_page_break()


def split_blocks(text: str):
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped or stripped == "---":
            i += 1
            continue
        if stripped.startswith("|") and stripped.endswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip())
                i += 1
            yield ("table", rows)
            continue
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            yield ("heading", level, stripped[level:].strip())
            i += 1
            continue
        m = re.match(r"^\s*[-*]\s+(.+)$", line)
        if m:
            yield ("bullet", m.group(1).strip())
            i += 1
            continue
        m = re.match(r"^\s*(\d+)\.\s+(.+)$", line)
        if m:
            yield ("number", m.group(1), m.group(2).strip())
            i += 1
            continue
        paragraph = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if (
                not nxt
                or nxt == "---"
                or nxt.startswith("#")
                or (nxt.startswith("|") and nxt.endswith("|"))
                or re.match(r"^[-*]\s+", nxt)
                or re.match(r"^\d+\.\s+", nxt)
            ):
                break
            paragraph.append(nxt)
            i += 1
        yield ("paragraph", " ".join(paragraph))


def parse_table_row(line: str):
    return [part.strip() for part in line.strip().strip("|").split("|")]


def add_markdown_table(doc: Document, rows: list[str]) -> None:
    parsed = [parse_table_row(row) for row in rows]
    if len(parsed) > 1 and all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in parsed[1]):
        parsed.pop(1)
    width = max(len(row) for row in parsed)
    parsed = [row + [""] * (width - len(row)) for row in parsed]
    table = doc.add_table(rows=len(parsed), cols=width)
    set_table_borders(table)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    header_tr_pr = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_tr_pr.append(repeat)
    for r_idx, values in enumerate(parsed):
        tr_pr = table.rows[r_idx]._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        for c_idx, value in enumerate(values):
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            set_cell_margins(cell, 65, 75, 65, 75)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if r_idx == 0:
                set_cell_shading(cell, BLUE)
            elif r_idx % 2 == 0:
                set_cell_shading(cell, PALE_BLUE)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, value)
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(7.5 if width >= 4 else 8.5)
                if r_idx == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(WHITE)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def render_footprint(doc: Document, source: Path) -> None:
    text = source.read_text(encoding="utf-8")
    for block in split_blocks(text):
        kind = block[0]
        if kind == "heading":
            _, level, title = block
            if level == 1:
                continue
            if title.lower() == "bảng thông tin cá nhân":
                continue
            style = "Heading 1" if level == 2 else "Heading 2" if level == 3 else "Heading 3"
            doc.add_paragraph(title, style=style)
        elif kind == "paragraph":
            p = doc.add_paragraph()
            add_inline(p, block[1])
        elif kind == "bullet":
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Cm(0.65)
            p.paragraph_format.first_line_indent = Cm(-0.35)
            add_inline(p, "• " + block[1])
        elif kind == "number":
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.left_indent = Cm(0.65)
            p.paragraph_format.first_line_indent = Cm(-0.45)
            add_inline(p, f"{block[1]}. {block[2]}")
        elif kind == "table":
            rows = block[1]
            header = " ".join(rows[:1]).lower()
            if "thông tin" in header and "nội dung" in header:
                continue
            add_markdown_table(doc, rows)


def add_confirmation(doc: Document, meta: dict) -> None:
    heading = doc.add_paragraph("INDIVIDUAL DECLARATION", style="Heading 1")
    heading.paragraph_format.page_break_before = True
    p = doc.add_paragraph()
    add_inline(
        p,
        "I confirm that this report reflects the work that I am responsible for explaining, calculating, "
        "demonstrating or reproducing in Finfolio 2.0. Referenced outputs may be collaborative team artifacts; "
        "my individual accountability is defined by the contribution map and repository evidence.",
    )
    table = doc.add_table(rows=4, cols=2)
    set_table_borders(table)
    fields = [
        ("Full name", meta["name"]),
        ("Student ID", meta["student_id"]),
        ("Signature", ""),
        ("Date", ""),
    ]
    for row, (label, value) in zip(table.rows, fields):
        row.cells[0].text = label
        row.cells[1].text = value
        set_cell_shading(row.cells[0], LIGHT_BLUE)
        for cell in row.cells:
            set_cell_margins(cell, 150, 120, 150, 120)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(10.5)
        row.cells[0].paragraphs[0].runs[0].bold = True


def set_update_fields(doc: Document) -> None:
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def build_one(slug: str, meta: dict, logo_path: Path) -> Path:
    doc = Document(REFERENCE)
    clear_document(doc)
    configure_styles(doc)
    configure_cover_section(doc)
    set_update_fields(doc)
    add_cover(doc, meta, logo_path)
    configure_body_section(doc)
    add_front_matter(doc, meta)
    render_footprint(doc, SOURCE_DIR / f"{slug}.md")
    add_confirmation(doc, meta)

    core = doc.core_properties
    core.title = f"Personal Project Report — {meta['name']}"
    core.subject = PROJECT
    core.author = meta["name"]
    core.keywords = "Finfolio 2.0, G09, individual footprint, NHA408E"

    output = OUTPUT_DIR / f"Finfolio_Personal_Report_{slug}_FINAL.docx"
    doc.save(output)
    return output


def main() -> None:
    if not REFERENCE.exists():
        raise FileNotFoundError(f"Reference document not found: {REFERENCE}")
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="finfolio-logo-") as tmp:
        logo_path = Path(tmp) / "ftu-logo.png"
        with zipfile.ZipFile(REFERENCE) as zf:
            logo_path.write_bytes(zf.read("word/media/image1.png"))
        outputs = [build_one(slug, meta, logo_path) for slug, meta in MEMBERS.items()]
    print("Created:")
    for output in outputs:
        print(f"- {output}")


if __name__ == "__main__":
    main()
