from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "output" / "slide_assets"
OUTPUT_FILE = OUTPUT_DIR / "finfolio_asset_class_attribution_mindmap_editable.pptx"

BG = "F7F4EA"
FOREST = "123D32"
GREEN = "2B6651"
SAGE = "E5EEE8"
SAGE_LIGHT = "F0F4EF"
GOLD = "C7A44B"
GOLD_LIGHT = "F5E8B4"
INK = "1D2925"
MUTED = "596861"
WHITE = "FFFFFF"


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def set_text(shape, text, size, color=INK, bold=False, align=PP_ALIGN.LEFT,
             font="Aptos", margin=0.08, valign=MSO_ANCHOR.MIDDLE):
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(0.03)
    frame.margin_bottom = Inches(0.03)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)
    paragraph.line_spacing = 1.0
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return shape


def add_text(slide, x, y, w, h, text, size, color=INK, bold=False,
             align=PP_ALIGN.LEFT, font="Aptos", valign=MSO_ANCHOR.MIDDLE):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    return set_text(box, text, size, color, bold, align, font, 0, valign)


def add_line(slide, x1, y1, x2, y2, color=GREEN, width=1.5):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    line.line.color.rgb = rgb(color)
    line.line.width = Pt(width)
    return line


def add_panel(slide, x, y, w, h, number, title, lines, highlight=None,
              body_size=12.5):
    panel = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    panel.fill.solid()
    panel.fill.fore_color.rgb = rgb(WHITE)
    panel.line.color.rgb = rgb(GREEN)
    panel.line.width = Pt(1.5)

    header_h = 0.42
    header = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x + 0.02), Inches(y + 0.02),
        Inches(w - 0.04), Inches(header_h)
    )
    header.fill.solid()
    header.fill.fore_color.rgb = rgb(SAGE)
    header.line.fill.background()

    badge = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(x + 0.13), Inches(y + 0.065), Inches(0.32), Inches(0.32)
    )
    badge.fill.solid()
    badge.fill.fore_color.rgb = rgb(GOLD)
    badge.line.fill.background()
    set_text(badge, str(number), 12, WHITE, True, PP_ALIGN.CENTER, margin=0)

    add_text(slide, x + 0.55, y + 0.04, w - 0.68, 0.36, title, 14.5,
             FOREST, True)

    body_y = y + 0.49
    body_h = h - 0.57
    if highlight is not None:
        row_h = body_h / len(lines)
        hi = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(x + 0.19), Inches(body_y + row_h * highlight + row_h * 0.40),
            Inches(w - 0.38), Inches(max(0.12, row_h * 0.75))
        )
        hi.fill.solid()
        hi.fill.fore_color.rgb = rgb(GOLD_LIGHT)
        hi.line.fill.background()

    body = slide.shapes.add_textbox(
        Inches(x + 0.22), Inches(body_y), Inches(w - 0.44), Inches(body_h)
    )
    frame = body.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = Inches(0.04)
    frame.margin_top = frame.margin_bottom = Inches(0)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for idx, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.space_before = paragraph.space_after = Pt(0)
        paragraph.line_spacing = 0.95
        run = paragraph.add_run()
        run.text = line
        run.font.name = "Aptos"
        run.font.size = Pt(body_size)
        run.font.bold = idx == highlight
        run.font.color.rgb = rgb(INK)
    return panel


def add_drill_box(slide, x, y, w, heading, value):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.62)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(WHITE)
    shape.line.color.rgb = rgb(GOLD)
    shape.line.width = Pt(1.25)
    add_text(slide, x + 0.08, y + 0.035, w - 0.16, 0.19, heading, 8.5,
             FOREST, True, PP_ALIGN.CENTER)
    value_box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x + 0.06), Inches(y + 0.27),
        Inches(w - 0.12), Inches(0.28)
    )
    value_box.fill.solid()
    value_box.fill.fore_color.rgb = rgb(GOLD_LIGHT)
    value_box.line.fill.background()
    set_text(value_box, value, 10.5, FOREST, True, PP_ALIGN.CENTER, margin=0.02)
    return shape


def add_chevron(slide, x, y):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.CHEVRON, Inches(x), Inches(y), Inches(0.28), Inches(0.28)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(GOLD)
    shape.line.fill.background()
    return shape


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = rgb(BG)

    add_text(slide, 0.4, 0.15, 12.53, 0.46,
             "HOW FINFOLIO GROUPS PORTFOLIO ATTRIBUTION", 25,
             FOREST, True, PP_ALIGN.CENTER, "Aptos Display")
    add_text(slide, 0.65, 0.64, 12.03, 0.3,
             "The same holding-level contributions can be aggregated through alternative, mutually exclusive views",
             12.5, GREEN, False, PP_ALIGN.CENTER)

    # Connectors are intentionally drawn before the nodes so their endpoints sit behind panels.
    left_x, left_w = 0.42, 3.88
    right_x, right_w = 9.03, 3.88
    center_x, center_y, center_w, center_h = 5.22, 2.22, 2.90, 2.34
    for y in (1.42, 2.92, 4.58):
        add_line(slide, left_x + left_w, y, center_x + 0.12, center_y + center_h / 2)
    for y in (1.42, 2.92, 4.58):
        add_line(slide, center_x + center_w - 0.12, center_y + center_h / 2,
                 right_x, y)

    add_panel(slide, left_x, 1.08, left_w, 0.75, 1, "BY INSTRUMENT",
              ["FPT.VN · HPG.VN · GLD"], body_size=12.5)
    add_panel(slide, left_x, 2.02, left_w, 2.03, 2, "BY ASSET CLASS",
              ["Equity", "Equity ETF", "Commodity Proxy", "Bond ETF",
               "Real Estate ETF", "Crypto"], highlight=2, body_size=11.4)
    add_panel(slide, left_x, 4.23, left_w, 1.25, 3, "BY EXPOSURE GROUP",
              ["Vietnam Equity · US Equity", "Gold · Silver · Oil · Copper",
               "Agriculture · Broad Commodities"], body_size=11.4)

    add_panel(slide, right_x, 1.08, right_w, 0.75, 4, "BY CURRENCY",
              ["VND · USD"], body_size=12.5)
    add_panel(slide, right_x, 2.02, right_w, 2.03, 5, "BY SECTOR",
              ["Corporate Sectors", "Precious Metals", "Energy",
               "Industrial / Base Metals", "Agriculture / Grains",
               "Broad Commodities"], highlight=1, body_size=11.2)
    add_panel(slide, right_x, 4.23, right_w, 1.72, 6, "BY INDUSTRY",
              ["Corporate Industries", "Monetary Metals", "Precious Industrial Metals",
               "Conventional Energy · Natural Gas", "Construction & Infrastructure Metals",
               "Grains & Oilseeds · Softs", "Commodity Basket"],
              highlight=1, body_size=9.7)

    center = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(center_x), Inches(center_y),
        Inches(center_w), Inches(center_h)
    )
    center.fill.solid()
    center.fill.fore_color.rgb = rgb(FOREST)
    center.line.color.rgb = rgb(GOLD)
    center.line.width = Pt(2.0)
    set_text(center, "PORTFOLIO\nATTRIBUTION\n\nSum of holding\ncontributions",
             18, WHITE, True, PP_ALIGN.CENTER, "Aptos Display", margin=0.12)

    add_text(slide, 0.55, 6.07, 12.23, 0.24,
             "COMMODITY DRILL-DOWN EXAMPLE", 10, GOLD, True, PP_ALIGN.CENTER)
    drill_y, drill_w = 6.32, 2.38
    drill_x = [0.60, 3.63, 6.66, 9.69]
    drill = [
        ("ASSET CLASS", "Commodity Proxy"),
        ("SECTOR", "Precious Metals"),
        ("INDUSTRY", "Monetary Metals"),
        ("INSTRUMENT", "GLD"),
    ]
    for idx, (heading, value) in enumerate(drill):
        add_drill_box(slide, drill_x[idx], drill_y, drill_w, heading, value)
        if idx < len(drill) - 1:
            add_chevron(slide, drill_x[idx] + drill_w + 0.35, drill_y + 0.17)

    footer = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.55), Inches(7.08), Inches(12.23), Inches(0.28)
    )
    footer.fill.solid()
    footer.fill.fore_color.rgb = rgb(SAGE)
    footer.line.fill.background()
    set_text(footer,
             "Choose one grouping view at a time. Every view reconciles to total portfolio return and risk.",
             10.5, FOREST, False, PP_ALIGN.CENTER, margin=0.03)

    # Useful metadata when the deck is handed between teammates.
    prs.core_properties.title = "Finfolio Portfolio Attribution Grouping Mind Map"
    prs.core_properties.subject = "Editable 16:9 slide asset"
    prs.core_properties.author = "Finfolio G09"
    prs.core_properties.comments = (
        "All diagram components are editable PowerPoint objects. "
        "Grouping views aggregate the same mutually exclusive holding-level contributions."
    )

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT_FILE)
    print(OUTPUT_FILE)


if __name__ == "__main__":
    build()
